import csv
import io
from datetime import datetime
from typing import Any, Optional

from app.core.timezone import format_ist_datetime, now_ist

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import os
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.pdfgen import canvas



def export_to_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([str(item) if item is not None else "" for item in row])
    return output.getvalue().encode("utf-8")


def export_to_excel(sheet_title: str, headers: list[str], rows: list[list[Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel max sheet title is 31 chars

    # Styling tokens matching UI spec: deep indigo #3C3489
    header_fill = PatternFill(start_color="3C3489", end_color="3C3489", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_side = Side(border_style="thin", color="E2E8F0")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Header Row
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    # Data Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append([str(v) if v is not None else "" for v in row_data])
        ws.row_dimensions[row_idx].height = 20
        use_alt = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border
            if use_alt:
                cell.fill = alt_fill

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class GeneralReportNumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 7)
        self.setFillColor(colors.HexColor("#64748B"))
        page_w, page_h = A4
        text = f"Global Agencies Godown Management  •  Page {self._pageNumber} of {page_count}"
        self.drawCentredString(page_w / 2.0, 14, text)
        self.restoreState()


def export_to_pdf(
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[list[Any]],
    col_widths: Optional[list[float]] = None,
) -> bytes:
    """
    Renders tabular godown reports calibrated for true A4 Portrait fit (595 x 842 pt).
    Computes proportional column widths to ensure the table never exceeds page margins.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=28,
        rightMargin=28,
        topMargin=26,
        bottomMargin=26,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=15,
        textColor=colors.HexColor("#1E1B4B"),
    )
    sub_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#64748B"),
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=8.5,
    )
    head_style = ParagraphStyle(
        "TableHead",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=colors.white,
    )

    story = [
        Paragraph(title, title_style),
        Paragraph(f"{subtitle}  •  Generated: {format_ist_datetime(now_ist())}", sub_style),
        Spacer(1, 8),
    ]

    # Proportional column width distribution to guarantee true A4 width fit (538 pt)
    content_width = 538.0
    has_real_records = any(
        r and any(str(cell).strip() not in ("", "-", "None", "No records found matching criteria") for cell in r)
        for r in rows
    )

    if not col_widths and headers:
        col_weights = []
        for c_idx in range(len(headers)):
            h_text = headers[c_idx]
            max_len = max(len(h_text), 6)
            if has_real_records:
                for r in rows[:100]:
                    if c_idx < len(r):
                        max_len = max(max_len, len(str(r[c_idx] or "")))
            col_weights.append(min(max_len, 32))
        tot_weight = sum(col_weights) or 1.0
        col_widths = [(w / tot_weight) * content_width for w in col_weights]

    # Convert headers to Paragraphs
    table_data = [[Paragraph(h, head_style) for h in headers]]
    table_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3C3489")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
    ]

    if not has_real_records:
        empty_msg = Paragraph('<font color="#64748B"><i>No records found matching criteria for this period</i></font>', cell_style)
        dummy_row = [empty_msg]
        for _ in range(len(headers) - 1):
            dummy_row.append(Paragraph("", cell_style))
        table_data.append(dummy_row)
        table_styles.append(("SPAN", (0, 1), (-1, 1)))
        table_styles.append(("ALIGN", (0, 1), (-1, 1), "CENTER"))
    else:
        for r in rows:
            row_items = []
            for c_idx in range(len(headers)):
                val = r[c_idx] if c_idx < len(r) and r[c_idx] is not None else ""
                row_items.append(Paragraph(str(val), cell_style))
            table_data.append(row_items)

    t = Table(table_data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle(table_styles))
    story.append(t)
    doc.build(story, canvasmaker=GeneralReportNumberedCanvas)
    return buf.getvalue()


class InwardRegisterNumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 7)
        self.setFillColor(colors.HexColor("#64748B"))
        page_w, page_h = A4
        text = f"Global Agencies Godown Management  •  Inward Serial Register  •  Page {self._pageNumber} of {page_count}"
        self.drawCentredString(page_w / 2.0, 14, text)
        self.restoreState()


def export_inward_register_to_pdf(
    title: str,
    subtitle: str,
    filter_desc: str,
    model_groups: list[dict[str, Any]],
    totals: dict[str, int],
) -> bytes:
    """
    Renders official Today Inward Serial Register grouped cleanly by model for A4 portrait print.
    Designed with compact margins and proportions so everyday inward operations fit on a single A4 sheet.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=28,
        rightMargin=28,
        topMargin=26,
        bottomMargin=26,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "RegTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=15,
        textColor=colors.HexColor("#1E1B4B"),
    )
    sub_style = ParagraphStyle(
        "RegSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#64748B"),
    )
    model_bar_left = ParagraphStyle(
        "ModelBarLeft",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )
    model_bar_right = ParagraphStyle(
        "ModelBarRight",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#FEF08A"),
        alignment=2,
    )
    head_style = ParagraphStyle(
        "RegTh",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#334155"),
        alignment=0,
    )
    cell_style = ParagraphStyle(
        "RegTd",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor("#0F172A"),
    )
    serial_cell = ParagraphStyle(
        "RegSn",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#1E1B4B"),
    )
    badge_stock_in = ParagraphStyle(
        "BadgeIn",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor("#047857"),
    )
    badge_return = ParagraphStyle(
        "BadgeRet",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor("#4338CA"),
    )
    badge_damaged = ParagraphStyle(
        "BadgeDam",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor("#B91C1C"),
    )

    story = []

    # 1. Header Banner
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 2))
    gen_time = format_ist_datetime(now_ist())
    meta_line = f"{subtitle}  •  Generated: {gen_time}"
    story.append(Paragraph(meta_line, sub_style))
    story.append(Spacer(1, 5))

    # 2. Stats summary bar (total width = 538 pt)
    total_models = totals.get("total_models", 0)
    total_serials = totals.get("total_serials", 0)
    stock_in = totals.get("stock_in", 0)
    returns = totals.get("returns", 0)
    damaged = totals.get("damaged", 0)

    summary_data = [
        [
            Paragraph(f"<b>Filter:</b> {filter_desc}", cell_style),
            Paragraph(f"<b>Models:</b> {total_models}", cell_style),
            Paragraph(f"<b>Total:</b> <font color='#3C3489'><b>{total_serials}</b></font>", cell_style),
            Paragraph(f"<b>Stock In:</b> <font color='#059669'><b>{stock_in}</b></font>", cell_style),
            Paragraph(f"<b>Returns:</b> <font color='#4338CA'><b>{returns}</b></font>", cell_style),
            Paragraph(f"<b>Damaged:</b> <font color='#DC2626'><b>{damaged}</b></font>", cell_style),
        ]
    ]
    summary_table = Table(summary_data, colWidths=[148, 78, 78, 78, 78, 78])
    summary_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ])
    )
    story.append(summary_table)
    story.append(Spacer(1, 6))

    if not model_groups:
        no_rec = ParagraphStyle("NoRec", parent=styles["Normal"], fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#64748B"), alignment=1)
        story.append(Spacer(1, 30))
        story.append(Paragraph("No inward serial numbers found matching criteria for today.", no_rec))
        doc.build(story, canvasmaker=InwardRegisterNumberedCanvas)
        return buf.getvalue()

    # Column widths total = 538 pt (matches A4 printable width)
    col_widths = [24, 145, 52, 67, 95, 55, 100]

    for grp in model_groups:
        brand = grp.get("brand") or "Unknown"
        model = grp.get("model") or "Model"
        pname = grp.get("product_name") or ""
        cat = grp.get("category") or ""
        serials = grp.get("serials") or []
        count = len(serials)

        # Model Header Bar (538 pt)
        left_txt = f"{brand.upper()} &nbsp;•&nbsp; <b>{model}</b> &nbsp; <font color='#E0E7FF' size=7.5>({pname}) &nbsp;[{cat}]</font>"
        right_txt = f"TOTAL: {count} UNIT{'S' if count != 1 else ''}"
        model_bar = Table(
            [[Paragraph(left_txt, model_bar_left), Paragraph(right_txt, model_bar_right)]],
            colWidths=[418, 120],
        )
        model_bar.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#3C3489")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ])
        )

        # Serials Table
        table_rows = [
            [
                Paragraph("#", head_style),
                Paragraph("Serial Number", head_style),
                Paragraph("Unit", head_style),
                Paragraph("Inward Type", head_style),
                Paragraph("Invoice / DC Ref", head_style),
                Paragraph("Time", head_style),
                Paragraph("Received By", head_style),
            ]
        ]

        for idx, s in enumerate(serials, start=1):
            itype = s.get("inward_type", "Stock In")
            if "Return" in itype:
                b_style = badge_return
            elif "Damaged" in itype:
                b_style = badge_damaged
            else:
                b_style = badge_stock_in

            table_rows.append([
                Paragraph(str(idx), cell_style),
                Paragraph(s.get("serial_number", "-"), serial_cell),
                Paragraph(s.get("unit_type", "-").capitalize() if s.get("unit_type") else "-", cell_style),
                Paragraph(itype, b_style),
                Paragraph(s.get("reference") or "-", cell_style),
                Paragraph(s.get("time") or "-", cell_style),
                Paragraph(s.get("received_by") or "-", cell_style),
            ])

        serials_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        serials_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ])
        )

        story.append(model_bar)
        story.append(serials_table)
        story.append(Spacer(1, 6))

    # Verification / Sign-off Footer (538 pt)
    sign_data = [
        [
            Paragraph(f"<b>Report Summary:</b> {total_serials} Total Serial(s) across {total_models} Model(s).", cell_style),
            Paragraph("Godown Inward Verifier Signature: ___________________________", cell_style),
        ]
    ]
    sign_table = Table(sign_data, colWidths=[288, 250])
    sign_table.setStyle(
        TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ])
    )
    story.append(Spacer(1, 4))
    story.append(sign_table)

    doc.build(story, canvasmaker=InwardRegisterNumberedCanvas)
    return buf.getvalue()



class BillNumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 7)
        self.setFillColor(colors.HexColor("#334155"))
        page_w, page_h = A4
        bill_num = getattr(self, "_bill_number", "")
        text = f"This is a Computer Generated Delivery Challan | Bill #{bill_num} | Page {self._pageNumber} of {page_count}"
        self.drawCentredString(page_w / 2.0, 15, text)
        self.restoreState()


def export_bill_to_pdf(bill: dict[str, Any]) -> bytes:
    """
    Renders the official Delivery Challan PDF matching frontend/src/utils/generateBillPdf.ts
    using the jsPDF generator script. Falls back to ReportLab if Node.js is unavailable.
    """
    import json
    import subprocess
    import tempfile

    script_path = os.path.join(os.path.dirname(__file__), "generate_bill_pdf.js")
    if os.path.exists(script_path):
        in_path = None
        out_path = None
        try:
            with tempfile.NamedTemporaryFile("w+", suffix=".json", delete=False, encoding="utf-8") as in_f:
                json.dump(bill, in_f)
                in_path = in_f.name

            with tempfile.NamedTemporaryFile("r+", suffix=".pdf", delete=False) as out_f:
                out_path = out_f.name

            cmd = ["node", script_path, in_path, out_path]
            proc = subprocess.run(cmd, capture_output=True, timeout=15)
            if proc.returncode == 0 and os.path.exists(out_path) and os.path.getsize(out_path) > 0:
                with open(out_path, "rb") as f:
                    pdf_data = f.read()
                try:
                    os.unlink(in_path)
                    os.unlink(out_path)
                except Exception:
                    pass
                return pdf_data
        except Exception:
            pass
        finally:
            for p in (in_path, out_path):
                if p and os.path.exists(p):
                    try:
                        os.unlink(p)
                    except Exception:
                        pass

    # Fallback to ReportLab implementation below
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=34,
        rightMargin=34,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    content_width = 527.0  # 595.27 - 68

    title_style = ParagraphStyle(
        "TallyTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        alignment=1,
        textColor=colors.black,
    )
    sub_title_style = ParagraphStyle(
        "TallySubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        alignment=1,
        textColor=colors.black,
    )
    bold_7 = ParagraphStyle("B7", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=7, leading=8.5, textColor=colors.black)
    bold_8 = ParagraphStyle("B8", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.black)
    bold_9 = ParagraphStyle("B9", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=11, textColor=colors.black)
    norm_7 = ParagraphStyle("N7", parent=styles["Normal"], fontName="Helvetica", fontSize=7.2, leading=9, textColor=colors.black)
    norm_75 = ParagraphStyle("N75", parent=styles["Normal"], fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=colors.black)
    norm_8 = ParagraphStyle("N8", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.black)
    courier_8 = ParagraphStyle("C8", parent=styles["Normal"], fontName="Courier", fontSize=8, leading=11.5, textColor=colors.black)

    story = []

    # 1. TOP HEADER BOX
    shop_name = bill.get("shop_name", "UNKNOWN SHOP")
    shop_city = bill.get("shop_city") or ""
    remarks = bill.get("remarks") or ""
    dest_str = f"Location: {shop_city}" if shop_city else ""
    rem_str = f"Remarks: {remarks}" if remarks else ""

    # Optional Logo
    logo_widget = None
    assets_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets")
    logo_file = os.path.join(assets_dir, "logo.png")
    if not os.path.exists(logo_file):
        logo_file = os.path.join("a:/GA STOCK/mobile/assets/images/logo.png")
    if os.path.exists(logo_file):
        try:
            logo_widget = Image(logo_file, width=28, height=28)
        except Exception:
            logo_widget = None

    if logo_widget:
        consignor_text = [
            Paragraph("<b>GLOBAL LOGISTICS</b>", ParagraphStyle("GL", parent=bold_9, fontSize=10, leading=12)),
            Paragraph("Central Warehouse & Distribution", norm_7),
            Paragraph("Stock Management & Automated Serial Register", norm_7),
        ]
        consignor_row = Table([[logo_widget, consignor_text]], colWidths=[34, 236])
        consignor_row.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        left_consignor = [
            Paragraph("<b>Consignor / Dispatcher:</b>", bold_7),
            Spacer(1, 2),
            consignor_row,
        ]
    else:
        left_consignor = [
            Paragraph("<b>Consignor / Dispatcher:</b>", bold_7),
            Spacer(1, 2),
            Paragraph("<b>GLOBAL LOGISTICS</b>", ParagraphStyle("GL", parent=bold_9, fontSize=10.5, leading=12.5)),
            Paragraph("Central Warehouse & Distribution", norm_7),
            Paragraph("Stock Management & Automated Serial Register", norm_7),
        ]

    left_consignee = [
        Paragraph("<b>Consignee / Buyer (Dispatch To):</b>", bold_7),
        Spacer(1, 2),
        Paragraph(f"<b>{shop_name}</b>", bold_9),
    ]
    if dest_str:
        left_consignee.append(Paragraph(dest_str, norm_7))
    if rem_str:
        left_consignee.append(Paragraph(rem_str, norm_7))

    left_table = Table([[left_consignor], [left_consignee]], colWidths=[280])
    left_table.setStyle(TableStyle([
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))

    bill_num = bill.get("bill_number", "-")
    txn_date = bill.get("transaction_date", "-")
    disp_name = bill.get("dispatched_by_name") or bill.get("dispatcher_name") or "-"
    vehicle = bill.get("delivery_reference") or "Direct / Local"
    destination = shop_city if shop_city else shop_name

    right_meta = [
        Paragraph(f"Challan / Bill No.:  <b>{bill_num}</b>", norm_75),
        Spacer(1, 2.5),
        Paragraph(f"Dated:  <b>{txn_date}</b>", norm_75),
        Spacer(1, 2.5),
        Paragraph(f"Dispatched By:  <b>{disp_name}</b>", norm_75),
        Spacer(1, 2.5),
        Paragraph(f"Vehicle / Ref:  <b>{vehicle}</b>", norm_75),
        Spacer(1, 2.5),
        Paragraph("Mode of Transport:  Road / Local Delivery", norm_75),
        Spacer(1, 2.5),
        Paragraph(f"Destination:  {destination}", norm_75),
    ]

    title_block = [
        Paragraph("DELIVERY CHALLAN", title_style),
        Spacer(1, 2),
        Paragraph("(CENTRAL WAREHOUSE OUTWARD DISPATCH / STOCK TRANSFER VOUCHER)", sub_title_style),
        Spacer(1, 3),
    ]

    header_table = Table([
        [Table([[title_block]], colWidths=[content_width])],
        [Table([[left_table, right_meta]], colWidths=[285, 242])],
    ], colWidths=[content_width])

    header_table.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.25, colors.black),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 8))

    # 2. TABLE 1: DESCRIPTION OF GOODS DISPATCHED
    story.append(Paragraph("<b>1. Description of Goods Dispatched</b>", bold_9))
    story.append(Spacer(1, 3))

    batches = bill.get("batches", [])
    model_head = [
        Paragraph("<b>SI No.</b>", ParagraphStyle("THC", parent=bold_8, alignment=1)),
        Paragraph("<b>Brand</b>", bold_8),
        Paragraph("<b>Model</b>", bold_8),
        Paragraph("<b>Description / Specification</b>", bold_8),
        Paragraph("<b>Quantity</b>", ParagraphStyle("THR", parent=bold_8, alignment=2)),
    ]
    model_rows = [model_head]
    total_qty = 0

    for idx, b in enumerate(batches, start=1):
        brand = b.get("brand") if isinstance(b, dict) else getattr(b, "brand", "-")
        model = b.get("model") if isinstance(b, dict) else getattr(b, "model", "-")
        p_name = b.get("product_name") if isinstance(b, dict) else getattr(b, "product_name", "-")
        qty = b.get("quantity") if isinstance(b, dict) else getattr(b, "quantity", 0)
        total_qty += qty

        model_rows.append([
            Paragraph(str(idx), ParagraphStyle("MC", parent=norm_8, alignment=1)),
            Paragraph(f"<b>{brand}</b>", bold_8),
            Paragraph(f"<b>{model}</b>", bold_8),
            Paragraph(str(p_name or "-"), norm_8),
            Paragraph(f"<b>{qty} Units</b>", ParagraphStyle("MR", parent=bold_8, alignment=2)),
        ])

    # Total row
    model_rows.append([
        "",
        Paragraph("<b>TOTAL</b>", bold_8),
        Paragraph(f"<b>{len(batches)} Model(s)</b>", bold_8),
        "",
        Paragraph(f"<b>{total_qty} Units</b>", ParagraphStyle("TR", parent=bold_8, alignment=2)),
    ])

    m_table = Table(model_rows, colWidths=[38, 90, 120, 189, 90], repeatRows=1)
    m_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F2F2F2")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(m_table)
    story.append(Spacer(1, 10))

    # 3. TABLE 2: SERIAL NUMBER REGISTER (GROUPED BY MODEL & 3 PER LINE)
    story.append(Paragraph("<b>2. Serial Number Register</b>", bold_9))
    story.append(Spacer(1, 3))

    serial_table_data = [
        [
            Paragraph("<b>Model Details & Serial Number Register</b>", bold_8),
            Paragraph("<b>Quantity</b>", ParagraphStyle("SHR", parent=bold_8, alignment=2)),
        ]
    ]

    for b_idx, b in enumerate(batches, start=1):
        brand = b.get("brand") if isinstance(b, dict) else getattr(b, "brand", "-")
        model = b.get("model") if isinstance(b, dict) else getattr(b, "model", "-")
        p_name = b.get("product_name") if isinstance(b, dict) else getattr(b, "product_name", "")
        desc_text = f"  |  {p_name}" if (p_name and p_name != f"{brand} {model}") else ""
        qty = b.get("quantity") if isinstance(b, dict) else getattr(b, "quantity", 0)
        lines = b.get("lines", []) if isinstance(b, dict) else getattr(b, "lines", [])

        # Model header subrow
        serial_table_data.append([
            Paragraph(f"<b>{b_idx}.  {brand} - {model}{desc_text}</b>", bold_8),
            Paragraph(f"<b>{qty} {'Unit' if qty == 1 else 'Units'}</b>", ParagraphStyle("SHR2", parent=bold_8, alignment=2)),
        ])

        # Serial numbers packed 3 per line
        if not lines:
            serial_table_data.append([
                Paragraph("<i>No scanned serial numbers recorded for this batch</i>", norm_75),
                "",
            ])
        else:
            serial_entries = []
            for l_idx, l in enumerate(lines, start=1):
                s_text = l.get("serial_text") if isinstance(l, dict) else getattr(l, "serial_text", "-")
                u_type = l.get("unit_type") if isinstance(l, dict) else getattr(l, "unit_type", None)
                u_tag = f" [{u_type.upper()}]" if (u_type and str(u_type).strip()) else ""
                serial_entries.append(f"{l_idx}.&nbsp;{s_text}{u_tag}")

            max_len = max((len(s.replace("&nbsp;", " ")) for s in serial_entries), default=0)
            items_per_line = 2 if max_len > 28 else 3
            col_width = 52 if items_per_line == 2 else 34

            formatted_lines = []
            for k in range(0, len(serial_entries), items_per_line):
                chunk = serial_entries[k:k+items_per_line]
                row_parts = []
                for c_idx, item_str in enumerate(chunk):
                    if c_idx == len(chunk) - 1:
                        row_parts.append(item_str)
                    else:
                        pad_len = max(2, col_width - len(item_str.replace("&nbsp;", " ")))
                        row_parts.append(item_str + ("&nbsp;" * pad_len))
                formatted_lines.append("".join(row_parts))

            body_html = "<br/>".join(formatted_lines)
            serial_table_data.append([
                Paragraph(body_html, courier_8),
                "",
            ])

    s_table = Table(serial_table_data, colWidths=[437, 90], repeatRows=1)
    s_table_style = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]

    row_cursor = 1
    for b in batches:
        # Model row
        s_table_style.append(("BACKGROUND", (0, row_cursor), (-1, row_cursor), colors.HexColor("#F2F2F2")))
        row_cursor += 1
        # Serial body row (span 2 columns)
        s_table_style.append(("SPAN", (0, row_cursor), (1, row_cursor)))
        s_table_style.append(("TOPPADDING", (0, row_cursor), (1, row_cursor), 5))
        s_table_style.append(("BOTTOMPADDING", (0, row_cursor), (1, row_cursor), 6))
        row_cursor += 1

    s_table.setStyle(TableStyle(s_table_style))
    story.append(s_table)
    story.append(Spacer(1, 10))

    # 4. DECLARATION & SIGNATURE FOOTER BOX
    left_sign = [
        Paragraph("<b>Declaration:</b>", bold_7),
        Spacer(1, 2),
        Paragraph(
            "We declare that this Delivery Challan shows the actual quantity of goods dispatched and all serial numbers listed have been physically scanned and verified prior to warehouse release.",
            norm_7,
        ),
        Spacer(1, 14),
        Paragraph("<b>Customer / Receiver's Signature & Seal</b>", bold_7),
    ]

    right_sign = [
        Paragraph("<b>for GLOBAL LOGISTICS</b>", bold_8),
        Paragraph("(Central Warehouse Dispatch Authority)", norm_7),
        Spacer(1, 18),
        Paragraph("<b>Authorised Signatory</b>", ParagraphStyle("AS", parent=bold_7, alignment=2)),
    ]

    sign_table = Table([[left_sign, right_sign]], colWidths=[310, 217])
    sign_table.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.25, colors.black),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    story.append(KeepTogether([sign_table]))

    def set_canvas_bill(canvas_obj, doc_obj):
        canvas_obj._bill_number = bill_num

    doc.build(story, canvasmaker=BillNumberedCanvas, onFirstPage=set_canvas_bill, onLaterPages=set_canvas_bill)
    return buf.getvalue()